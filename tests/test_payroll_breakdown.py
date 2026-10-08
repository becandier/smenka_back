# tests/test_payroll_breakdown.py
"""Фича payroll_breakdown: категории начислений, разбивка отчёта payroll и Excel."""

import uuid
from datetime import UTC, datetime
from decimal import ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.app.core.security import hash_password
from src.app.models.adjustment import PayrollAdjustmentCategory
from src.app.models.audit_log import AuditLog
from src.app.models.member_rate import OrganizationMemberRate, RateType
from src.app.models.organization import MemberRole, Organization, OrganizationMember
from src.app.models.penalty import Penalty
from src.app.models.shift import Shift, ShiftStatus
from src.app.models.shift_overtime_request import OvertimeRequestStatus, ShiftOvertimeRequest
from src.app.models.user import User
from src.app.services import adjustment as adjustment_service
from src.app.services.adjustment import AdjustmentError

RATE_EFF_JAN = datetime(2026, 1, 1, tzinfo=UTC)
JUNE = {"date_from": "2026-06-01T00:00:00Z", "date_to": "2026-06-30T23:59:59Z"}


def _data(resp: Any) -> Any:
    return resp.json()["data"]


def _err(resp: Any) -> str:
    return resp.json()["error"]["code"]


# --- fixtures ------------------------------------------------------------------
async def _make_user(db_session: AsyncSession, email: str, name: str) -> User:
    user = User(
        id=uuid.uuid4(),
        email=email,
        password_hash=hash_password("Test1234"),
        name=name,
        is_verified=True,
    )
    db_session.add(user)
    await db_session.commit()
    return user


async def _login(client: AsyncClient, email: str) -> dict[str, str]:
    resp = await client.post("/api/v1/auth/login", json={"email": email, "password": "Test1234"})
    return {"Authorization": f"Bearer {resp.json()['data']['access_token']}"}


@pytest.fixture
async def owner(db_session: AsyncSession) -> User:
    return await _make_user(db_session, "pb_owner@example.com", "Owner")


@pytest.fixture
async def owner_headers(owner: User, client: AsyncClient) -> dict[str, str]:
    return await _login(client, "pb_owner@example.com")


@pytest.fixture
async def emp2_user(db_session: AsyncSession) -> User:
    return await _make_user(db_session, "pb_emp2@example.com", "Anna Second")


@pytest.fixture
async def emp2_headers(emp2_user: User, client: AsyncClient) -> dict[str, str]:
    return await _login(client, "pb_emp2@example.com")


@pytest.fixture
async def org(db_session: AsyncSession, owner: User) -> Organization:
    organization = Organization(name="Breakdown Org", owner_id=owner.id)
    db_session.add(organization)
    await db_session.commit()
    return organization


@pytest.fixture
async def other_org(db_session: AsyncSession, owner: User) -> Organization:
    organization = Organization(name="Other Org", owner_id=owner.id)
    db_session.add(organization)
    await db_session.commit()
    return organization


@pytest.fixture
async def employee_member(
    db_session: AsyncSession, org: Organization, verified_user: User
) -> OrganizationMember:
    member = OrganizationMember(
        organization_id=org.id, user_id=verified_user.id, role=MemberRole.employee
    )
    db_session.add(member)
    await db_session.commit()
    return member


@pytest.fixture
async def emp2_member(
    db_session: AsyncSession, org: Organization, emp2_user: User
) -> OrganizationMember:
    member = OrganizationMember(
        organization_id=org.id, user_id=emp2_user.id, role=MemberRole.employee
    )
    db_session.add(member)
    await db_session.commit()
    return member


async def _make_finished_shift(
    db_session: AsyncSession,
    user_id: uuid.UUID,
    org_id: uuid.UUID,
    started_at: datetime,
    finished_at: datetime,
) -> Shift:
    shift = Shift(
        user_id=user_id,
        organization_id=org_id,
        started_at=started_at,
        finished_at=finished_at,
        status=ShiftStatus.finished,
    )
    db_session.add(shift)
    await db_session.commit()
    return shift


async def _approve_overtime(db_session: AsyncSession, shift_id: uuid.UUID, minutes: int) -> None:
    db_session.add(
        ShiftOvertimeRequest(
            shift_id=shift_id,
            minutes=minutes,
            comment="Задержался",
            status=OvertimeRequestStatus.approved,
        )
    )
    await db_session.commit()


async def _make_rate(
    db_session: AsyncSession,
    member_id: uuid.UUID,
    amount: int,
    rate_type: RateType = RateType.hourly,
) -> None:
    db_session.add(
        OrganizationMemberRate(
            member_id=member_id,
            rate_amount_minor=amount,
            rate_type=rate_type,
            currency="RUB",
            effective_from=RATE_EFF_JAN,
        )
    )
    await db_session.commit()


async def _make_penalty(
    db_session: AsyncSession,
    org_id: uuid.UUID,
    member_id: uuid.UUID,
    created_by: uuid.UUID,
    amount: int,
    occurred_at: datetime,
    *,
    reason: str = "Опоздание",
    shift_id: uuid.UUID | None = None,
    is_deleted: bool = False,
) -> Penalty:
    penalty = Penalty(
        organization_id=org_id,
        member_id=member_id,
        shift_id=shift_id,
        reason=reason,
        amount_minor=amount,
        occurred_at=occurred_at,
        created_by_user_id=created_by,
        is_deleted=is_deleted,
    )
    db_session.add(penalty)
    await db_session.commit()
    return penalty


async def _create_category(
    client: AsyncClient, headers: dict[str, str], org_id: uuid.UUID, name: str
) -> Any:
    return await client.post(
        f"/api/v1/organizations/{org_id}/adjustment-categories",
        headers=headers,
        json={"name": name},
    )


async def _category_id(
    client: AsyncClient, headers: dict[str, str], org_id: uuid.UUID, name: str
) -> str:
    resp = await _create_category(client, headers, org_id, name)
    assert resp.status_code == 201, resp.text
    return str(_data(resp)["id"])


async def _create_adjustment(
    client: AsyncClient, headers: dict[str, str], org_id: uuid.UUID, **body: Any
) -> Any:
    body.setdefault("reason", "Корректировка")
    body.setdefault("occurred_at", "2026-06-15T10:00:00Z")
    return await client.post(
        f"/api/v1/organizations/{org_id}/adjustments", headers=headers, json=body
    )


async def _adjustment(
    client: AsyncClient, headers: dict[str, str], org_id: uuid.UUID, **body: Any
) -> dict[str, Any]:
    resp = await _create_adjustment(client, headers, org_id, **body)
    assert resp.status_code == 201, resp.text
    data: dict[str, Any] = _data(resp)
    return data


# --- CRUD категорий -------------------------------------------------------------
async def test_category_crud_and_sorting(client, owner_headers, org, employee_member):
    base = f"/api/v1/organizations/{org.id}/adjustment-categories"
    resp = await _create_category(client, owner_headers, org.id, "  премия  ")
    assert resp.status_code == 201, resp.text
    bonus = _data(resp)
    assert bonus["name"] == "премия"  # trim
    assert bonus["organization_id"] == str(org.id)
    assert bonus["is_deleted"] is False
    assert bonus["deleted_at"] is None
    assert bonus["adjustments_count"] == 0

    await _category_id(client, owner_headers, org.id, "Компенсация")
    await _category_id(client, owner_headers, org.id, "Аванс")

    resp = await client.get(base, headers=owner_headers)
    assert resp.status_code == 200
    names = [c["name"] for c in _data(resp)["items"]]
    assert names == ["Аванс", "Компенсация", "премия"]  # lower(name)

    # adjustments_count — только неотменённые начисления
    a1 = await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=100,
        category_id=bonus["id"],
    )
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=200,
        category_id=bonus["id"],
    )
    await client.delete(
        f"/api/v1/organizations/{org.id}/adjustments/{a1['id']}", headers=owner_headers
    )

    resp = await client.patch(
        f"{base}/{bonus['id']}", headers=owner_headers, json={"name": "Премия"}
    )
    assert resp.status_code == 200, resp.text
    assert _data(resp)["name"] == "Премия"
    assert _data(resp)["adjustments_count"] == 1

    resp = await client.delete(f"{base}/{bonus['id']}", headers=owner_headers)
    assert resp.status_code == 200, resp.text
    assert _data(resp) == {"deleted": True}

    resp = await client.get(base, headers=owner_headers)
    assert [c["name"] for c in _data(resp)["items"]] == ["Аванс", "Компенсация"]

    resp = await client.get(base, headers=owner_headers, params={"include_deleted": "true"})
    deleted = next(c for c in _data(resp)["items"] if c["id"] == bonus["id"])
    assert deleted["is_deleted"] is True
    assert deleted["deleted_at"] is not None
    assert deleted["adjustments_count"] == 1


async def test_category_duplicate_case_insensitive_409(client, owner_headers, org):
    await _category_id(client, owner_headers, org.id, "Премия")
    resp = await _create_category(client, owner_headers, org.id, " ПРЕМИЯ ")
    assert resp.status_code == 409
    assert _err(resp) == "ADJUSTMENT_CATEGORY_DUPLICATE"

    other = await _category_id(client, owner_headers, org.id, "Компенсация")
    resp = await client.patch(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{other}",
        headers=owner_headers,
        json={"name": "премия"},
    )
    assert resp.status_code == 409
    assert _err(resp) == "ADJUSTMENT_CATEGORY_DUPLICATE"

    # переименование в то же имя с другим регистром у самой себя — не дубль
    resp = await client.patch(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{other}",
        headers=owner_headers,
        json={"name": "КОМПЕНСАЦИЯ"},
    )
    assert resp.status_code == 200, resp.text


async def test_category_recreate_after_delete(client, owner_headers, org):
    first = await _category_id(client, owner_headers, org.id, "Премия")
    await client.delete(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{first}", headers=owner_headers
    )
    resp = await _create_category(client, owner_headers, org.id, "премия")
    assert resp.status_code == 201, resp.text
    assert _data(resp)["id"] != first


async def test_category_same_name_in_other_org_allowed(client, owner_headers, org, other_org):
    await _category_id(client, owner_headers, org.id, "Премия")
    resp = await _create_category(client, owner_headers, other_org.id, "Премия")
    assert resp.status_code == 201, resp.text


async def test_category_race_unique_index_409(db_session, org, owner):
    """Гонка: проверка «имя свободно» прошла, но живая запись уже вставлена
    параллельно → нарушение частичного уникального индекса → 409, не 500."""
    db_session.add(
        PayrollAdjustmentCategory(
            organization_id=org.id, name="Премия", created_by_user_id=owner.id
        )
    )
    await db_session.commit()

    original = adjustment_service._ensure_category_name_free

    async def _skip_check(*args: Any, **kwargs: Any) -> None:
        return None

    adjustment_service._ensure_category_name_free = _skip_check  # type: ignore[assignment]
    try:
        with pytest.raises(AdjustmentError) as exc_info:
            await adjustment_service.create_category(db_session, org.id, owner.id, name="ПРЕМИЯ")
    finally:
        adjustment_service._ensure_category_name_free = original  # type: ignore[assignment]
        await db_session.rollback()
    assert exc_info.value.code == "ADJUSTMENT_CATEGORY_DUPLICATE"
    assert exc_info.value.status_code == 409


async def test_category_not_found_cases(client, owner_headers, org, other_org):
    base = f"/api/v1/organizations/{org.id}/adjustment-categories"
    foreign = await _category_id(client, owner_headers, other_org.id, "Чужая")
    for resp in (
        await client.patch(f"{base}/{foreign}", headers=owner_headers, json={"name": "X"}),
        await client.delete(f"{base}/{foreign}", headers=owner_headers),
        await client.patch(f"{base}/{uuid.uuid4()}", headers=owner_headers, json={"name": "X"}),
    ):
        assert resp.status_code == 404
        assert _err(resp) == "ADJUSTMENT_CATEGORY_NOT_FOUND"

    own = await _category_id(client, owner_headers, org.id, "Своя")
    assert (await client.delete(f"{base}/{own}", headers=owner_headers)).status_code == 200
    resp = await client.delete(f"{base}/{own}", headers=owner_headers)
    assert resp.status_code == 404
    assert _err(resp) == "ADJUSTMENT_CATEGORY_NOT_FOUND"
    resp = await client.patch(f"{base}/{own}", headers=owner_headers, json={"name": "Новая"})
    assert resp.status_code == 404
    assert _err(resp) == "ADJUSTMENT_CATEGORY_NOT_FOUND"


@pytest.mark.parametrize("name", ["", "   ", "x" * 101])
async def test_category_name_validation_422(client, owner_headers, org, name):
    resp = await _create_category(client, owner_headers, org.id, name)
    assert resp.status_code == 422
    assert _err(resp) == "VALIDATION_ERROR"


async def test_category_name_max_length_ok(client, owner_headers, org):
    resp = await _create_category(client, owner_headers, org.id, "я" * 100)
    assert resp.status_code == 201, resp.text


async def test_category_rbac(
    client, owner_headers, auth_headers, super_admin_headers, org, employee_member
):
    cat = await _category_id(client, owner_headers, org.id, "Премия")
    base = f"/api/v1/organizations/{org.id}/adjustment-categories"
    for headers in (auth_headers, super_admin_headers):
        for resp in (
            await client.get(base, headers=headers),
            await _create_category(client, headers, org.id, "Y"),
            await client.patch(f"{base}/{cat}", headers=headers, json={"name": "Z"}),
            await client.delete(f"{base}/{cat}", headers=headers),
        ):
            assert resp.status_code == 403, resp.text
            assert _err(resp) == "FORBIDDEN"


async def test_category_admin_member_allowed(
    client, db_session, owner_headers, org, emp2_user, emp2_headers
):
    db_session.add(
        OrganizationMember(organization_id=org.id, user_id=emp2_user.id, role=MemberRole.admin)
    )
    await db_session.commit()
    resp = await _create_category(client, emp2_headers, org.id, "Премия")
    assert resp.status_code == 201, resp.text


# --- Категория в ручных начислениях ---------------------------------------------
async def test_adjustment_create_with_category(
    client, owner_headers, org, employee_member, db_session
):
    cat = await _category_id(client, owner_headers, org.id, "Премия")
    data = await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=5000,
        category_id=cat,
    )
    assert data["category_id"] == cat
    assert data["category_name"] == "Премия"

    plain = await _adjustment(
        client, owner_headers, org.id, member_id=str(employee_member.id), amount_minor=100
    )
    assert plain["category_id"] is None
    assert plain["category_name"] is None

    audit = (
        await db_session.execute(
            select(AuditLog).where(
                AuditLog.action == "adjustment.create",
                AuditLog.resource_id == uuid.UUID(data["id"]),
            )
        )
    ).scalar_one()
    assert audit.summary["category_id"] == cat


async def test_adjustment_create_foreign_or_deleted_category_404(
    client, owner_headers, org, other_org, employee_member
):
    foreign = await _category_id(client, owner_headers, other_org.id, "Чужая")
    deleted = await _category_id(client, owner_headers, org.id, "Удалённая")
    await client.delete(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{deleted}", headers=owner_headers
    )
    for cat in (foreign, deleted, str(uuid.uuid4())):
        resp = await _create_adjustment(
            client,
            owner_headers,
            org.id,
            member_id=str(employee_member.id),
            amount_minor=100,
            category_id=cat,
        )
        assert resp.status_code == 404
        assert _err(resp) == "ADJUSTMENT_CATEGORY_NOT_FOUND"


async def test_adjustment_update_category_set_reset_keep(
    client, owner_headers, org, employee_member, db_session
):
    cat_a = await _category_id(client, owner_headers, org.id, "Премия")
    cat_b = await _category_id(client, owner_headers, org.id, "Компенсация")
    adj = await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=100,
        category_id=cat_a,
    )
    url = f"/api/v1/organizations/{org.id}/adjustments/{adj['id']}"

    # не передано — категория не меняется
    resp = await client.patch(url, headers=owner_headers, json={"reason": "Новая причина"})
    assert resp.status_code == 200, resp.text
    assert _data(resp)["category_id"] == cat_a

    # другая категория
    resp = await client.patch(url, headers=owner_headers, json={"category_id": cat_b})
    assert _data(resp)["category_id"] == cat_b
    assert _data(resp)["category_name"] == "Компенсация"
    audit = (
        (
            await db_session.execute(
                select(AuditLog)
                .where(AuditLog.action == "adjustment.update")
                .order_by(AuditLog.created_at.desc())
            )
        )
        .scalars()
        .first()
    )
    assert audit is not None
    assert audit.summary["changed"]["category_id"] == {"from": cat_a, "to": cat_b}

    # явный null — сброс
    resp = await client.patch(url, headers=owner_headers, json={"category_id": None})
    assert resp.status_code == 200, resp.text
    assert _data(resp)["category_id"] is None
    assert _data(resp)["category_name"] is None


async def test_adjustment_on_deleted_category_stays_valid(
    client, owner_headers, org, employee_member
):
    cat = await _category_id(client, owner_headers, org.id, "Премия")
    other = await _category_id(client, owner_headers, org.id, "Другая")
    adj = await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=100,
        category_id=cat,
    )
    await client.delete(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{cat}", headers=owner_headers
    )
    await client.patch(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{other}",
        headers=owner_headers,
        json={"name": "Другая2"},
    )
    url = f"/api/v1/organizations/{org.id}/adjustments/{adj['id']}"

    # правка другого поля и даже повтор той же (удалённой) категории — валидны
    resp = await client.patch(url, headers=owner_headers, json={"amount_minor": 300})
    assert resp.status_code == 200, resp.text
    assert _data(resp)["category_id"] == cat
    assert _data(resp)["category_name"] == "Премия"
    resp = await client.patch(url, headers=owner_headers, json={"category_id": cat})
    assert resp.status_code == 200, resp.text

    # список показывает имя удалённой категории
    resp = await client.get(f"/api/v1/organizations/{org.id}/adjustments", headers=owner_headers)
    item = next(i for i in _data(resp)["items"] if i["id"] == adj["id"])
    assert item["category_name"] == "Премия"

    # назначить удалённую категорию другому начислению — нельзя
    adj2 = await _adjustment(
        client, owner_headers, org.id, member_id=str(employee_member.id), amount_minor=100
    )
    resp = await client.patch(
        f"/api/v1/organizations/{org.id}/adjustments/{adj2['id']}",
        headers=owner_headers,
        json={"category_id": cat},
    )
    assert resp.status_code == 404
    assert _err(resp) == "ADJUSTMENT_CATEGORY_NOT_FOUND"


async def test_adjustment_rename_category_reflected(client, owner_headers, org, employee_member):
    cat = await _category_id(client, owner_headers, org.id, "Премия")
    adj = await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=100,
        category_id=cat,
    )
    await client.patch(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{cat}",
        headers=owner_headers,
        json={"name": "Бонус"},
    )
    resp = await client.get(f"/api/v1/organizations/{org.id}/adjustments", headers=owner_headers)
    item = next(i for i in _data(resp)["items"] if i["id"] == adj["id"])
    assert item["category_name"] == "Бонус"


async def test_adjustment_list_category_filter(client, owner_headers, org, employee_member):
    cat_a = await _category_id(client, owner_headers, org.id, "Премия")
    cat_b = await _category_id(client, owner_headers, org.id, "Компенсация")
    member = str(employee_member.id)
    a = await _adjustment(
        client, owner_headers, org.id, member_id=member, amount_minor=1, category_id=cat_a
    )
    b = await _adjustment(
        client, owner_headers, org.id, member_id=member, amount_minor=2, category_id=cat_b
    )
    n = await _adjustment(client, owner_headers, org.id, member_id=member, amount_minor=3)
    url = f"/api/v1/organizations/{org.id}/adjustments"

    resp = await client.get(url, headers=owner_headers, params={"category_id": cat_a})
    assert [i["id"] for i in _data(resp)["items"]] == [a["id"]]
    assert _data(resp)["total"] == 1
    resp = await client.get(url, headers=owner_headers, params={"category_id": cat_b})
    assert [i["id"] for i in _data(resp)["items"]] == [b["id"]]
    for token in ("none", "NONE"):
        resp = await client.get(url, headers=owner_headers, params={"category_id": token})
        assert [i["id"] for i in _data(resp)["items"]] == [n["id"]]
    resp = await client.get(url, headers=owner_headers, params={"category_id": "garbage"})
    assert resp.status_code == 422
    assert _err(resp) == "VALIDATION_ERROR"


async def test_my_adjustments_include_category(
    client, owner_headers, auth_headers, org, employee_member
):
    cat = await _category_id(client, owner_headers, org.id, "Премия")
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=100,
        category_id=cat,
    )
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(employee_member.id),
        amount_minor=-50,
        occurred_at="2026-06-14T10:00:00Z",
    )
    resp = await client.get(f"/api/v1/organizations/{org.id}/my-adjustments", headers=auth_headers)
    assert resp.status_code == 200, resp.text
    items = _data(resp)["items"]
    assert items[0]["category_id"] == cat
    assert items[0]["category_name"] == "Премия"
    assert items[1]["category_id"] is None
    assert items[1]["category_name"] is None


# --- Отчёт payroll: разбивка ------------------------------------------------------
def _amount(rate: int, seconds: int) -> int:
    exact = Decimal(seconds) * rate / Decimal(3600)
    return int(exact.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


HOURLY_RATE = 10001  # некруглая ставка — ловит расхождения округления


async def _seed_report(
    client: AsyncClient,
    db_session: AsyncSession,
    owner_headers: dict[str, str],
    org: Organization,
    owner: User,
    employee_member: OrganizationMember,
    emp2_member: OrganizationMember,
    verified_user: User,
    emp2_user: User,
) -> dict[str, Any]:
    """Test User — hourly с переработкой на двух сменах в разные недели;
    Anna Second — per_shift с переработкой. Начисления в 2 категориях + без
    категории, удалённая категория, отменённое начисление, штрафы."""
    await _make_rate(db_session, employee_member.id, HOURLY_RATE)
    await _make_rate(db_session, emp2_member.id, 500000, RateType.per_shift)

    s1 = await _make_finished_shift(
        db_session,
        verified_user.id,
        org.id,
        datetime(2026, 6, 3, 9, 0, 0, tzinfo=UTC),
        datetime(2026, 6, 3, 10, 0, 7, tzinfo=UTC),  # 3607 c
    )
    await _approve_overtime(db_session, s1.id, 13)
    s2 = await _make_finished_shift(
        db_session,
        verified_user.id,
        org.id,
        datetime(2026, 6, 20, 10, 0, tzinfo=UTC),
        datetime(2026, 6, 20, 11, 30, tzinfo=UTC),  # 5400 c
    )
    await _approve_overtime(db_session, s2.id, 7)
    s3 = await _make_finished_shift(
        db_session,
        emp2_user.id,
        org.id,
        datetime(2026, 6, 5, 9, 0, tzinfo=UTC),
        datetime(2026, 6, 5, 17, 0, tzinfo=UTC),
    )
    await _approve_overtime(db_session, s3.id, 30)

    bonus = await _category_id(client, owner_headers, org.id, "премия")
    advance = await _category_id(client, owner_headers, org.id, "Аванс")
    member = str(employee_member.id)
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=5000,
        category_id=bonus,
        reason="Премия за план",
        comment="Июнь",
        occurred_at="2026-06-10T10:00:00Z",
    )
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=3000,
        category_id=bonus,
        shift_id=str(s1.id),
        occurred_at=None,
    )
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=-2000,
        category_id=advance,
        occurred_at="2026-06-15T22:30:00Z",
    )
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=-1000,
        occurred_at="2026-06-12T10:00:00Z",
    )
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=700,
        occurred_at="2026-06-11T10:00:00Z",
    )
    cancelled = await _adjustment(
        client, owner_headers, org.id, member_id=member, amount_minor=99999, category_id=bonus
    )
    await client.delete(
        f"/api/v1/organizations/{org.id}/adjustments/{cancelled['id']}", headers=owner_headers
    )
    # вне периода — не попадает
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=777,
        category_id=bonus,
        occurred_at="2026-07-02T10:00:00Z",
    )
    # удалённая категория у Anna — остаётся в отчёте под своим именем
    gone = await _category_id(client, owner_headers, org.id, "Форма")
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(emp2_member.id),
        amount_minor=-400,
        category_id=gone,
    )
    await client.delete(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{gone}", headers=owner_headers
    )
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=str(emp2_member.id),
        amount_minor=100,
        category_id=bonus,
    )

    await _make_penalty(
        db_session,
        org.id,
        employee_member.id,
        owner.id,
        1500,
        datetime(2026, 6, 4, 9, 0, tzinfo=UTC),
        shift_id=s1.id,
    )
    await _make_penalty(
        db_session,
        org.id,
        employee_member.id,
        owner.id,
        9999,
        datetime(2026, 6, 4, 9, 0, tzinfo=UTC),
        is_deleted=True,
    )
    await _make_penalty(
        db_session,
        org.id,
        emp2_member.id,
        owner.id,
        250,
        datetime(2026, 6, 6, 9, 0, tzinfo=UTC),
        reason="Форма",
    )

    base1 = _amount(HOURLY_RATE, 3607) + _amount(HOURLY_RATE, 5400)
    gross1 = _amount(HOURLY_RATE, 3607 + 13 * 60) + _amount(HOURLY_RATE, 5400 + 7 * 60)
    return {"bonus": bonus, "advance": advance, "gone": gone, "base1": base1, "gross1": gross1}


def _assert_gross_split(container: dict[str, Any]) -> None:
    assert (
        container["base_amount_minor"] + container["overtime_amount_minor"]
        == container["gross_amount_minor"]
    )


def _assert_adjustment_invariants(container: dict[str, Any]) -> None:
    assert container["adjustment_accrual_minor"] >= 0
    assert container["adjustment_deduction_minor"] >= 0
    assert (
        container["adjustment_accrual_minor"] - container["adjustment_deduction_minor"]
        == container["adjustment_amount_minor"]
    )
    by_cat = container["adjustments_by_category"]
    assert sum(c["amount_minor"] for c in by_cat) == container["adjustment_amount_minor"]
    assert sum(c["count"] for c in by_cat) == container["adjustments_count"]
    for c in by_cat:
        assert c["accrual_minor"] - c["deduction_minor"] == c["amount_minor"]


@pytest.mark.parametrize("granularity", ["none", "day", "week", "month"])
async def test_payroll_breakdown_invariants(
    client,
    db_session,
    owner_headers,
    owner,
    org,
    employee_member,
    emp2_member,
    verified_user,
    emp2_user,
    granularity,
):
    seeded = await _seed_report(
        client,
        db_session,
        owner_headers,
        org,
        owner,
        employee_member,
        emp2_member,
        verified_user,
        emp2_user,
    )
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll",
        headers=owner_headers,
        params={**JUNE, "granularity": granularity},
    )
    assert resp.status_code == 200, resp.text
    data = _data(resp)
    items = {i["user_name"]: i for i in data["items"]}
    test_user, anna = items["Test User"], items["Anna Second"]

    # hourly с переработкой: точное значение base/overtime по сменам
    assert test_user["gross_amount_minor"] == seeded["gross1"]
    assert test_user["base_amount_minor"] == seeded["base1"]
    assert test_user["overtime_amount_minor"] == seeded["gross1"] - seeded["base1"]
    assert test_user["overtime_amount_minor"] > 0
    # per_shift: переработка не оплачивается отдельно
    assert anna["gross_amount_minor"] == 500000
    assert anna["base_amount_minor"] == 500000
    assert anna["overtime_amount_minor"] == 0

    for container in [*data["items"], data["totals"]]:
        _assert_gross_split(container)
        _assert_adjustment_invariants(container)
    if granularity != "none":
        for item in data["items"]:
            for bucket in item["breakdown"]:
                _assert_gross_split(bucket)
            assert (
                sum(b["base_amount_minor"] for b in item["breakdown"]) == item["base_amount_minor"]
            )
            assert (
                sum(b["overtime_amount_minor"] for b in item["breakdown"])
                == item["overtime_amount_minor"]
            )

    assert test_user["adjustment_amount_minor"] == 5700
    assert test_user["adjustment_accrual_minor"] == 8700
    assert test_user["adjustment_deduction_minor"] == 3000
    assert test_user["adjustments_count"] == 5
    assert test_user["adjustments_by_category"] == [
        {
            "category_id": seeded["advance"],
            "category_name": "Аванс",
            "category_is_deleted": False,
            "amount_minor": -2000,
            "accrual_minor": 0,
            "deduction_minor": 2000,
            "count": 1,
        },
        {
            "category_id": seeded["bonus"],
            "category_name": "премия",
            "category_is_deleted": False,
            "amount_minor": 8000,
            "accrual_minor": 8000,
            "deduction_minor": 0,
            "count": 2,
        },
        {
            "category_id": None,
            "category_name": None,
            "category_is_deleted": False,
            "amount_minor": -300,
            "accrual_minor": 700,
            "deduction_minor": 1000,
            "count": 2,
        },
    ]
    assert test_user["penalty_amount_minor"] == 1500
    assert test_user["net_amount_minor"] == seeded["gross1"] - 1500 + 5700

    assert anna["adjustments_by_category"] == [
        {
            "category_id": seeded["bonus"],
            "category_name": "премия",
            "category_is_deleted": False,
            "amount_minor": 100,
            "accrual_minor": 100,
            "deduction_minor": 0,
            "count": 1,
        },
        {
            "category_id": seeded["gone"],
            "category_name": "Форма",
            "category_is_deleted": True,
            "amount_minor": -400,
            "accrual_minor": 0,
            "deduction_minor": 400,
            "count": 1,
        },
    ]

    totals = data["totals"]
    assert totals["base_amount_minor"] == seeded["base1"] + 500000
    assert totals["adjustment_accrual_minor"] == 8800
    assert totals["adjustment_deduction_minor"] == 3400
    assert [
        (c["category_name"], c["category_is_deleted"], c["amount_minor"], c["count"])
        for c in totals["adjustments_by_category"]
    ] == [
        ("Аванс", False, -2000, 1),
        ("премия", False, 8100, 3),
        ("Форма", True, -400, 1),
        (None, False, -300, 2),
    ]


async def test_payroll_breakdown_include_adjustments_false(
    client,
    db_session,
    owner_headers,
    owner,
    org,
    employee_member,
    emp2_member,
    verified_user,
    emp2_user,
):
    await _seed_report(
        client,
        db_session,
        owner_headers,
        org,
        owner,
        employee_member,
        emp2_member,
        verified_user,
        emp2_user,
    )
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll",
        headers=owner_headers,
        params={**JUNE, "include_adjustments": "false"},
    )
    data = _data(resp)
    for container in [*data["items"], data["totals"]]:
        assert container["adjustment_amount_minor"] == 0
        assert container["adjustment_accrual_minor"] == 0
        assert container["adjustment_deduction_minor"] == 0
        assert container["adjustments_count"] == 0
        assert container["adjustments_by_category"] == []
        _assert_gross_split(container)


async def test_payroll_breakdown_user_filter_and_empty(
    client,
    db_session,
    owner_headers,
    owner,
    org,
    employee_member,
    emp2_member,
    verified_user,
    emp2_user,
):
    seeded = await _seed_report(
        client,
        db_session,
        owner_headers,
        org,
        owner,
        employee_member,
        emp2_member,
        verified_user,
        emp2_user,
    )
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll",
        headers=owner_headers,
        params={**JUNE, "user_ids": str(emp2_user.id)},
    )
    data = _data(resp)
    assert [i["user_name"] for i in data["items"]] == ["Anna Second"]
    assert [c["category_id"] for c in data["totals"]["adjustments_by_category"]] == [
        seeded["bonus"],
        seeded["gone"],
    ]
    assert data["totals"]["adjustment_amount_minor"] == -300

    # период без начислений — пустой список категорий
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll",
        headers=owner_headers,
        params={"date_from": "2026-05-01T00:00:00Z", "date_to": "2026-05-31T23:59:59Z"},
    )
    assert _data(resp)["totals"]["adjustments_by_category"] == []
    assert _data(resp)["items"] == []


# --- Excel -------------------------------------------------------------------------
def _sheet_rows(wb: Any, title: str) -> list[tuple[Any, ...]]:
    return list(wb[title].iter_rows(values_only=True))


def _summary(wb: Any) -> tuple[list[str], dict[str, tuple[Any, ...]]]:
    rows = _sheet_rows(wb, "Сводка")
    header = list(next(r for r in rows if r and r[0] == "Сотрудник"))
    body = {str(r[0]): r for r in rows[rows.index(tuple(header)) + 1 :] if r and r[0] is not None}
    return header, body


async def test_payroll_export_breakdown_sheets(
    client,
    db_session,
    owner_headers,
    owner,
    org,
    employee_member,
    emp2_member,
    verified_user,
    emp2_user,
):
    seeded = await _seed_report(
        client,
        db_session,
        owner_headers,
        org,
        owner,
        employee_member,
        emp2_member,
        verified_user,
        emp2_user,
    )
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll/export",
        headers=owner_headers,
        params={**JUNE, "tz": "Europe/Moscow"},
    )
    assert resp.status_code == 200, resp.text
    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Сводка", "Детализация", "Начисления и удержания", "Штрафы"]

    header, body = _summary(wb)
    assert header == [
        "Сотрудник",
        "Часы",
        "Смены",
        "Начислено, ₽",
        "в т.ч. за время, ₽",
        "в т.ч. переработка, ₽",
        "Штраф, ₽",
        "Доплаты, ₽",
        "Удержания, ₽",
        "Аванс, ₽",
        "премия, ₽",
        "Форма (удалена), ₽",
        "Без категории, ₽",
        "К выплате, ₽",
        "Без ставки (смен)",
        "Без ставки (часов)",
        "Переработка, ч",
        "По графику, ч",
        "По графику, ₽",
        "Разница, ₽",
        "Опозданий",
        "Опоздания, мин",
    ]
    col = {name: idx for idx, name in enumerate(header)}
    assert list(body) == ["Anna Second", "Test User", "ИТОГО"]  # как items
    tu = body["Test User"]
    assert tu[col["Начислено, ₽"]] == seeded["gross1"] / 100
    assert tu[col["в т.ч. за время, ₽"]] == seeded["base1"] / 100
    assert tu[col["Доплаты, ₽"]] == 87.0
    assert tu[col["Удержания, ₽"]] == -30.0
    assert tu[col["Аванс, ₽"]] == -20.0
    assert tu[col["премия, ₽"]] == 80.0
    assert tu[col["Форма (удалена), ₽"]] == 0
    assert tu[col["Без категории, ₽"]] == -3.0
    anna = body["Anna Second"]
    assert anna[col["в т.ч. переработка, ₽"]] == 0
    assert anna[col["Форма (удалена), ₽"]] == -4.0

    # К выплате = Начислено − Штраф + Доплаты + Удержания; Начислено = время + переработка
    for row in body.values():
        assert row[col["К выплате, ₽"]] == pytest.approx(
            row[col["Начислено, ₽"]]
            - row[col["Штраф, ₽"]]
            + row[col["Доплаты, ₽"]]
            + row[col["Удержания, ₽"]]
        )
        assert row[col["Начислено, ₽"]] == pytest.approx(
            row[col["в т.ч. за время, ₽"]] + row[col["в т.ч. переработка, ₽"]]
        )
        cats = sum(
            row[col[h]]
            for h in ("Аванс, ₽", "премия, ₽", "Форма (удалена), ₽", "Без категории, ₽")
        )
        assert cats == pytest.approx(row[col["Доплаты, ₽"]] + row[col["Удержания, ₽"]])
    # ИТОГО — по всем денежным колонкам, включая категорийные
    for name in (
        "Начислено, ₽",
        "в т.ч. за время, ₽",
        "в т.ч. переработка, ₽",
        "Штраф, ₽",
        "Доплаты, ₽",
        "Удержания, ₽",
        "Аванс, ₽",
        "премия, ₽",
        "Форма (удалена), ₽",
        "Без категории, ₽",
        "К выплате, ₽",
    ):
        assert body["ИТОГО"][col[name]] == pytest.approx(
            body["Test User"][col[name]] + body["Anna Second"][col[name]]
        )

    # Детализация: без вводящих в заблуждение колонок, с разбивкой gross
    detail = _sheet_rows(wb, "Детализация")
    assert list(detail[0]) == [
        "Сотрудник",
        "Дата",
        "Часы",
        "Смены",
        "Начислено, ₽",
        "в т.ч. за время, ₽",
        "в т.ч. переработка, ₽",
        "Без ставки (часов)",
        "По графику, ч",
        "По графику, ₽",
        "Разница, ₽",
    ]
    for r in detail[1:]:
        assert r[4] == pytest.approx(r[5] + r[6])

    # Лист начислений: те же строки, что в агрегате
    adj_rows = _sheet_rows(wb, "Начисления и удержания")
    assert list(adj_rows[0]) == [
        "Сотрудник",
        "Дата",
        "Категория",
        "Причина",
        "Комментарий",
        "Сумма, ₽",
        "Смена",
        "Кто внёс",
    ]
    adj_body = adj_rows[1:-1]
    assert adj_rows[-1][0] == "ИТОГО"
    assert [r[0] for r in adj_body] == ["Anna Second"] * 2 + ["Test User"] * 5
    for name in ("Test User", "Anna Second"):
        user_sum = sum(r[5] for r in adj_body if r[0] == name)
        assert user_sum == pytest.approx(
            body[name][col["Доплаты, ₽"]] + body[name][col["Удержания, ₽"]]
        )
    assert adj_rows[-1][5] == pytest.approx(sum(r[5] for r in adj_body))
    tu_rows = [r for r in adj_body if r[0] == "Test User"]
    # сортировка по дате; дата в tz отчёта (22:30Z 15.06 → 16.06 по Москве)
    assert [r[1] for r in tu_rows] == [
        "03.06.2026",
        "10.06.2026",
        "11.06.2026",
        "12.06.2026",
        "16.06.2026",
    ]
    shift_row = tu_rows[0]
    assert shift_row[2] == "премия"
    assert shift_row[6] == "03.06.2026 12:00"  # начало смены 09:00Z по Москве
    assert shift_row[7] == "Owner"
    assert tu_rows[1][3] == "Премия за план"
    assert tu_rows[1][4] == "Июнь"
    assert tu_rows[1][6] in ("", None)
    assert tu_rows[2][2] == "Без категории"
    assert sorted(r[2] for r in adj_body if r[0] == "Anna Second") == ["Форма (удалена)", "премия"]

    # Лист штрафов
    pen_rows = _sheet_rows(wb, "Штрафы")
    assert list(pen_rows[0]) == [
        "Сотрудник",
        "Дата",
        "Причина",
        "Комментарий",
        "Сумма, ₽",
        "Смена",
        "Кто назначил",
    ]
    pen_body = pen_rows[1:-1]
    assert len(pen_body) == 2  # отменённый штраф не попал
    for name in ("Test User", "Anna Second"):
        assert sum(r[4] for r in pen_body if r[0] == name) == pytest.approx(
            body[name][col["Штраф, ₽"]]
        )
    assert pen_rows[-1][0] == "ИТОГО"
    assert pen_rows[-1][4] == pytest.approx(17.5)


async def test_payroll_export_without_adjustments_and_penalties(
    client,
    db_session,
    owner_headers,
    owner,
    org,
    employee_member,
    emp2_member,
    verified_user,
    emp2_user,
):
    await _seed_report(
        client,
        db_session,
        owner_headers,
        org,
        owner,
        employee_member,
        emp2_member,
        verified_user,
        emp2_user,
    )
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll/export",
        headers=owner_headers,
        params={**JUNE, "include_adjustments": "false", "include_penalties": "false"},
    )
    wb = load_workbook(BytesIO(resp.content))
    assert wb.sheetnames == ["Сводка", "Детализация"]
    header, _ = _summary(wb)
    assert "Доплаты, ₽" not in header
    assert "Удержания, ₽" not in header
    assert not any(
        h.endswith(", ₽") and h.startswith(("Аванс", "премия", "Без кат")) for h in header
    )


async def test_payroll_export_user_filter_rows_match(
    client,
    db_session,
    owner_headers,
    owner,
    org,
    employee_member,
    emp2_member,
    verified_user,
    emp2_user,
):
    await _seed_report(
        client,
        db_session,
        owner_headers,
        org,
        owner,
        employee_member,
        emp2_member,
        verified_user,
        emp2_user,
    )
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll/export",
        headers=owner_headers,
        params={**JUNE, "user_ids": str(verified_user.id)},
    )
    wb = load_workbook(BytesIO(resp.content))
    header, body = _summary(wb)
    assert list(body) == ["Test User", "ИТОГО"]
    assert "Форма (удалена), ₽" not in header  # категория только у отфильтрованной Anna
    assert {r[0] for r in _sheet_rows(wb, "Начисления и удержания")[1:-1]} == {"Test User"}
    assert {r[0] for r in _sheet_rows(wb, "Штрафы")[1:-1]} == {"Test User"}


async def test_payroll_export_only_missing_rate_rows_match(
    client,
    db_session,
    owner_headers,
    owner,
    org,
    employee_member,
    emp2_member,
    verified_user,
    emp2_user,
):
    """only_missing_rate сужает items: Anna (со ставкой, без штрафов) остаётся только
    благодаря начислению; при include_adjustments=false она выпадает — строки листов
    строго по сотрудникам из «Сводки»."""
    await _make_rate(db_session, emp2_member.id, 10000)
    await _make_finished_shift(
        db_session,
        verified_user.id,
        org.id,
        datetime(2026, 6, 3, 9, 0, tzinfo=UTC),
        datetime(2026, 6, 3, 10, 0, tzinfo=UTC),
    )  # Test User без ставки → остаётся
    await _make_finished_shift(
        db_session,
        emp2_user.id,
        org.id,
        datetime(2026, 6, 3, 9, 0, tzinfo=UTC),
        datetime(2026, 6, 3, 10, 0, tzinfo=UTC),
    )
    await _make_penalty(
        db_session,
        org.id,
        employee_member.id,
        owner.id,
        300,
        datetime(2026, 6, 4, 9, 0, tzinfo=UTC),
    )
    await _adjustment(
        client, owner_headers, org.id, member_id=str(emp2_member.id), amount_minor=100
    )

    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll/export",
        headers=owner_headers,
        params={**JUNE, "only_missing_rate": "true"},
    )
    wb = load_workbook(BytesIO(resp.content))
    _, body = _summary(wb)
    assert list(body) == ["Anna Second", "Test User", "ИТОГО"]
    assert [r[0] for r in _sheet_rows(wb, "Начисления и удержания")[1:-1]] == ["Anna Second"]
    assert [r[0] for r in _sheet_rows(wb, "Штрафы")[1:-1]] == ["Test User"]

    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll/export",
        headers=owner_headers,
        params={**JUNE, "only_missing_rate": "true", "include_adjustments": "false"},
    )
    wb = load_workbook(BytesIO(resp.content))
    _, body = _summary(wb)
    assert list(body) == ["Test User", "ИТОГО"]
    assert [r[0] for r in _sheet_rows(wb, "Штрафы")[1:-1]] == ["Test User"]


async def test_adjustment_list_category_filter_checks_access_first(
    client, auth_headers, org, employee_member
):
    resp = await client.get(
        f"/api/v1/organizations/{org.id}/adjustments",
        headers=auth_headers,
        params={"category_id": "garbage"},
    )
    assert resp.status_code == 403
    assert _err(resp) == "FORBIDDEN"


async def test_category_audit_log(client, owner_headers, owner, org, db_session):
    base = f"/api/v1/organizations/{org.id}/adjustment-categories"
    cat = await _category_id(client, owner_headers, org.id, "Премия")
    # переименование в то же имя — не изменение, аудита нет
    await client.patch(f"{base}/{cat}", headers=owner_headers, json={"name": "Премия"})
    await client.patch(f"{base}/{cat}", headers=owner_headers, json={"name": "Бонус"})
    await client.delete(f"{base}/{cat}", headers=owner_headers)

    logs = (
        (
            await db_session.execute(
                select(AuditLog)
                .where(AuditLog.resource_id == uuid.UUID(cat))
                .order_by(AuditLog.created_at)
            )
        )
        .scalars()
        .all()
    )
    assert [log.action for log in logs] == [
        "adjustment_category.create",
        "adjustment_category.update",
        "adjustment_category.delete",
    ]
    assert all(log.resource_type == "adjustment_category" for log in logs)
    assert all(log.organization_id == org.id for log in logs)
    assert all(log.actor_user_id == owner.id for log in logs)
    assert logs[0].summary == {"name": "Премия"}
    assert logs[1].summary == {"changed": {"name": {"from": "Премия", "to": "Бонус"}}}
    assert logs[2].summary == {"name": "Бонус"}


async def test_payroll_deleted_and_live_category_same_name(
    client, db_session, owner_headers, org, employee_member, verified_user
):
    """Удалённая «Премия» и новая живая «премия»: две записи, живая раньше
    удалённой, в Excel удалённая — с суффиксом « (удалена)»."""
    await _make_rate(db_session, employee_member.id, 18000)
    member = str(employee_member.id)
    old = await _category_id(client, owner_headers, org.id, "Премия")
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=100,
        category_id=old,
        occurred_at="2026-06-02T10:00:00Z",
    )
    await client.delete(
        f"/api/v1/organizations/{org.id}/adjustment-categories/{old}", headers=owner_headers
    )
    new = await _category_id(client, owner_headers, org.id, "премия")
    await _adjustment(
        client,
        owner_headers,
        org.id,
        member_id=member,
        amount_minor=200,
        category_id=new,
        occurred_at="2026-06-03T10:00:00Z",
    )

    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll", headers=owner_headers, params=JUNE
    )
    by_cat = _data(resp)["totals"]["adjustments_by_category"]
    assert [(c["category_id"], c["category_is_deleted"], c["amount_minor"]) for c in by_cat] == [
        (new, False, 200),
        (old, True, 100),
    ]

    resp = await client.get(
        f"/api/v1/organizations/{org.id}/payroll/export", headers=owner_headers, params=JUNE
    )
    wb = load_workbook(BytesIO(resp.content))
    header, body = _summary(wb)
    assert "премия, ₽" in header
    assert "Премия (удалена), ₽" in header
    assert header.index("премия, ₽") < header.index("Премия (удалена), ₽")
    assert body["Test User"][header.index("Премия (удалена), ₽")] == 1.0
    labels = [r[2] for r in _sheet_rows(wb, "Начисления и удержания")[1:-1]]
    assert labels == ["Премия (удалена)", "премия"]
